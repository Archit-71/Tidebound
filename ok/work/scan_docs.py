from docx import Document
from pathlib import Path
root=Path(r'C:\Users\Archit\Documents\Codex\2026-10-14\ok\outputs')
files=['Tidebound_Project_Report.docx','Tidebound_Technical_Architecture_and_Design.docx','Tidebound_Combined_Project_and_Technical_Document.docx']
terms=['developer','github','super pirate','super-pirate','clear code','repository','repo name','upstream']
for name in files:
 p=root/name
 d=Document(p)
 print('\n###',name)
 for i,para in enumerate(d.paragraphs):
  text=para.text
  if any(t in text.lower() for t in terms): print('P',i,repr(text))
 for ti,table in enumerate(d.tables):
  for ri,row in enumerate(table.rows):
   for ci,cell in enumerate(row.cells):
    text=cell.text
    if any(t in text.lower() for t in terms): print('T',ti,ri,ci,repr(text))
