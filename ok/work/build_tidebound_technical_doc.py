from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'C:\Users\Archit\Documents\Codex\2026-10-14\ok\outputs')
WORK = Path(r'C:\Users\Archit\Documents\Codex\2026-10-14\ok\work')
OUT.mkdir(parents=True, exist_ok=True)
WORK.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / 'Tidebound_Technical_Architecture_and_Design.docx'

NAVY = '#14253D'
INK = '#1E2A36'
BLUE = '#2E6F95'
TEAL = '#3A9D9A'
GOLD = '#C99A3B'
PALE = '#EEF3F7'
MID = '#D4E0E8'
WHITE = '#FFFFFF'
GRAY = '#586876'

def font_path():
    for p in [Path(r'C:\Windows\Fonts\arial.ttf'), Path(r'C:\Windows\Fonts\segoeui.ttf')]:
        if p.exists():
            return str(p)
    return None

FONT = font_path()

def pil_font(size, bold=False):
    if FONT:
        p = FONT
        if bold:
            candidate = Path(r'C:\Windows\Fonts\arialbd.ttf')
            if candidate.exists(): p = str(candidate)
        return ImageFont.truetype(p, size)
    return ImageFont.load_default()

def hexrgb(value):
    value = value.lstrip('#')
    return tuple(int(value[i:i+2], 16) for i in (0, 2, 4))

def rounded(draw, box, fill, outline=None, radius=18, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)

def wrap_words(draw, text, font, max_width):
    lines, line = [], ''
    for word in text.split():
        trial = f'{line} {word}'.strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            line = trial
        else:
            if line: lines.append(line)
            line = word
    if line: lines.append(line)
    return lines

def diagram(name, title, nodes, edges, note):
    # nodes: {id: (x,y,w,h,label,kind)} in a 1600 x 900 drawing plane.
    scale = 1.4
    W, H = int(1600*scale), int(900*scale)
    im = Image.new('RGB', (W, H), hexrgb('#F7FAFC'))
    d = ImageDraw.Draw(im)
    title_font = pil_font(int(36*scale), True)
    node_font = pil_font(int(21*scale), True)
    note_font = pil_font(int(15*scale))
    d.text((int(60*scale), int(36*scale)), title, fill=hexrgb(NAVY), font=title_font)
    # arrows are drawn behind the node cards
    for a, b, label in edges:
        ax, ay, aw, ah, *_ = nodes[a]
        bx, by, bw, bh, *_ = nodes[b]
        start = (int((ax + aw/2)*scale), int((ay + ah/2)*scale))
        end = (int((bx + bw/2)*scale), int((by + bh/2)*scale))
        dx, dy = end[0]-start[0], end[1]-start[1]
        if abs(dx) > abs(dy):
            start = (int((ax + (aw if dx > 0 else 0))*scale), int((ay+ah/2)*scale))
            end = (int((bx + (0 if dx > 0 else bw))*scale), int((by+bh/2)*scale))
        else:
            start = (int((ax+aw/2)*scale), int((ay + (ah if dy > 0 else 0))*scale))
            end = (int((bx+bw/2)*scale), int((by + (0 if dy > 0 else bh))*scale))
        d.line((start, end), fill=hexrgb('#74889A'), width=int(4*scale))
        angle = __import__('math').atan2(end[1]-start[1], end[0]-start[0])
        size = 13*scale
        left = (end[0]-size*__import__('math').cos(angle-0.48), end[1]-size*__import__('math').sin(angle-0.48))
        right = (end[0]-size*__import__('math').cos(angle+0.48), end[1]-size*__import__('math').sin(angle+0.48))
        d.polygon([end, left, right], fill=hexrgb('#74889A'))
        # Flow is shown by arrow direction; explanatory conditions belong in
        # node labels and captions so they do not collide with connector lines.
    for key, (x,y,w,h,label,kind) in nodes.items():
        fill = {'primary': NAVY, 'data': '#EAF4F5', 'external': '#F3EEDF', 'ui': '#E8F0F7', 'danger': '#F6E8E5'}.get(kind, WHITE)
        outline = {'primary': NAVY, 'data': TEAL, 'external': GOLD, 'ui': BLUE, 'danger': '#B95C54'}.get(kind, MID)
        box = (int(x*scale), int(y*scale), int((x+w)*scale), int((y+h)*scale))
        rounded(d, box, hexrgb(fill), hexrgb(outline), int(20*scale), int(3*scale))
        lines = wrap_words(d, label, node_font, int((w-28)*scale))
        line_h = int(30*scale)
        total = len(lines)*line_h
        cy = (box[1]+box[3]-total)//2
        for line in lines:
            bb = d.textbbox((0,0), line, font=node_font)
            tx = (box[0]+box[2]-(bb[2]-bb[0]))//2
            d.text((tx,cy), line, fill=hexrgb(WHITE if kind=='primary' else INK), font=node_font)
            cy += line_h
    d.text((int(60*scale), int(850*scale)), note, fill=hexrgb(GRAY), font=note_font)
    path = WORK / f'{name}.png'
    im.save(path, dpi=(180,180))
    return path

DIAGRAMS = {}

DIAGRAMS['context'] = diagram('diagram_context', 'System context', {
 'player':(70,350,250,135,'Player','external'),
 'game':(470,310,300,210,'Tidebound game loop','primary'),
 'assets':(900,130,280,145,'Graphics audio and fonts','external'),
 'maps':(900,350,280,145,'TMX maps via PyTMX','external'),
 'display':(900,600,280,145,'Pygame display and mixer','ui'),
 'python':(1280,350,250,145,'Python runtime','data')}, [
 ('player','game','input'),('assets','game','load'),('maps','game','load'),('game','display','render / play'),('python','game','executes')],
 'The game runs locally. It has no network service, account system, or remote data dependency.')

DIAGRAMS['architecture'] = diagram('diagram_architecture', 'Runtime component architecture', {
 'game':(620,70,340,145,'Game\nmain.py','primary'),
 'data':(120,300,280,135,'Data\nhealth coins levels','data'),
 'ui':(480,300,280,135,'UI\nhearts and coins','ui'),
 'level':(850,300,280,135,'Level\nTMX stage','primary'),
 'world':(1220,300,280,135,'Overworld\nroute map','primary'),
 'player':(460,570,280,135,'Player\nmovement and combat','ui'),
 'sprites':(820,570,280,135,'Sprites and groups\nrendering and camera','data'),
 'enemies':(1180,570,280,135,'Enemies\nTooth Shell Pearl','danger')}, [
 ('game','data','owns'),('game','ui','owns'),('game','level','creates'),('game','world','creates'),('level','player','creates'),('level','sprites','uses'),('level','enemies','creates'),('data','ui','updates')],
 'Game is the composition root; Level and Overworld are mutually exclusive active stages.')

DIAGRAMS['loop'] = diagram('diagram_loop', 'Frame lifecycle', {
 'start':(80,350,240,120,'Start frame\nclock tick','primary'),
 'events':(390,350,250,120,'Process events\nquit or menu input','ui'),
 'update':(720,350,250,120,'Update stage\ninput timers physics','primary'),
 'collide':(1030,350,250,120,'Resolve contacts\nitems damage exits','danger'),
 'draw':(1330,350,200,120,'Draw world UI\nflip display','ui'),
 'over':(720,650,250,120,'Game over\nhealth at zero\nfreeze stage','danger'),
 'retry':(1030,650,250,120,'Retry same level\nEnter or R\nrestore session','data')}, [
 ('start','events',''),('events','update','active'),('update','collide',''),('collide','draw',''),('events','over','health is zero'),('over','retry','Enter or R'),('retry','start','next frame')],
 'The loop is capped at 60 frames per second; simulation receives elapsed seconds as dt.')

DIAGRAMS['states'] = diagram('diagram_states', 'Stage and session states', {
 'boot':(80,350,230,120,'Initialize\nassets and maps','primary'),
 'over':(400,130,260,130,'Overworld\nselect node','ui'),
 'level':(790,130,260,130,'Level\nplay stage','primary'),
 'respawn':(790,540,260,130,'Respawn same level\nlose one health','data'),
 'gameover':(1190,350,280,140,'Game over menu\nhealth at zero','danger'),
 'exit':(1190,650,230,100,'Quit process','external')}, [
 ('boot','over','start at node 0'),('over','level','Enter'),('level','over','flag unlock'),('level','respawn','fall below map'),('respawn','level','health remains'),('level','gameover','health <= 0'),('gameover','level','Enter / R retry'),('gameover','exit','Esc')],
 'A stage transition replaces current_stage. Only one gameplay stage is active at a time.')

DIAGRAMS['map_build'] = diagram('diagram_map_build', 'TMX level construction', {
 'tmx':(60,350,240,130,'TMX map\nlayers and objects','external'),
 'meta':(360,100,250,130,'Data layer\nbackground unlock limits','data'),
 'tiles':(360,310,250,130,'Tile layers\nBG Terrain FG Platforms','ui'),
 'objects':(360,530,250,130,'Object layers\nplayer scenery hazards','ui'),
 'groups':(740,310,250,130,'Sprite groups\ncollision damage items','primary'),
 'entities':(1080,180,250,130,'Entity instances\nplayer enemies objects','primary'),
 'stage':(1080,480,250,130,'Active Level\nupdate and draw','data'),
 'game':(1400,350,160,130,'Game loop','primary')}, [
 ('tmx','meta','read'),('tmx','tiles','iterate'),('tmx','objects','iterate'),('tiles','groups','assign'),('objects','entities','instantiate'),('entities','stage','join'),('groups','stage','register'),('stage','game','run')],
 'The layer names and custom properties in each TMX map form a content contract with Level.setup.')

DIAGRAMS['collision'] = diagram('diagram_collision', 'Player update and contact pipeline', {
 'input':(70,300,230,130,'Read keyboard\nset direction','ui'),
 'move':(360,300,230,130,'Apply movement\nusing dt','primary'),
 'solid':(650,120,240,130,'Solid collision\nterrain and crates','danger'),
 'platform':(650,310,240,130,'One way collision\nplatform skip timer','data'),
 'contact':(650,520,240,130,'Contact checks\nfloor and walls','ui'),
 'stage':(1010,300,230,130,'Stage checks\nhazards items attacks','primary'),
 'health':(1360,180,200,130,'Health timer\ninvulnerability','danger'),
 'render':(1360,450,200,130,'Animation and\nrender frame','ui')}, [
 ('input','move',''),('move','solid',''),('move','platform',''),('solid','contact',''),('platform','contact',''),('contact','stage',''),('stage','health','damage'),('stage','render',''),('health','render','')],
 'Level.run performs stage-level interactions after all sprites update for the frame.')

DIAGRAMS['model'] = diagram('diagram_model', 'Session data and UI relationships', {
 'game':(650,60,300,130,'Game\nactive session owner','primary'),
 'data':(650,320,300,150,'Data\nhealth coins current level\nunlocked level','data'),
 'ui':(110,570,300,140,'UI\nheart sprites coin display','ui'),
 'level':(650,600,300,140,'Level\nreads and changes Data','primary'),
 'over':(1200,570,300,140,'Overworld\nreads level progress','ui'),
 'items':(120,320,300,140,'Item\nupdates coins or health','external'),
 'player':(1200,320,300,140,'Player\nreceives damage','danger')}, [
 ('game','data','creates'),('game','level','passes'),('game','over','passes'),('data','ui','setter callback'),('items','data','collect'),('player','data','damage'),('level','data','shared reference'),('over','data','shared reference')],
 'Data is in-memory only. It is recreated when retrying and is not written to disk.')

DIAGRAMS['world'] = diagram('diagram_world', 'Overworld route selection', {
 'nodes':(70,330,260,150,'TMX Nodes\nstage and route properties','external'),
 'paths':(420,120,270,140,'TMX Paths\nordered point lists','external'),
 'sprite':(420,450,270,140,'Node sprites\nlocked route check','data'),
 'input':(790,330,260,150,'Arrow input\nmove icon along route','ui'),
 'current':(1130,120,280,140,'Current node\nupdated by overlap','primary'),
 'enter':(1130,480,280,140,'Enter key\nsets current_level','primary'),
 'game':(1430,330,150,150,'Game\nloads Level','primary')}, [
 ('nodes','sprite','instantiate'),('paths','input','route geometry'),('sprite','input','movement options'),('input','current','icon arrives'),('current','enter','selection'),('enter','game','switch stage')],
 'Available paths are constrained by Data.unlocked_level; the map itself defines adjacency.')

DIAGRAMS['render'] = diagram('diagram_render', 'World rendering and camera', {
 'target':(60,350,230,130,'Player or icon\nworld position','primary'),
 'camera':(370,350,230,130,'Camera offset\nconstrained to map','data'),
 'groups':(690,350,240,130,'Sprite groups\nordered by z layer','ui'),
 'world':(1030,160,240,130,'World blits\nbackground and sprites','primary'),
 'ui':(1030,540,240,130,'UI blits\nhearts and coins','ui'),
 'display':(1370,350,200,130,'Display surface\nflip to screen','external')}, [
 ('target','camera','center view'),('camera','groups','offset'),('groups','world','draw'),('world','display',''),('ui','display','')],
 'Level scenes use AllSprites and camera offset. The overworld uses WorldSprites and its own camera.')

DIAGRAMS['death'] = diagram('diagram_death', 'Fall, depletion, and retry sequence', {
 'fall':(70,330,250,140,'Player falls below\nlevel boundary','danger'),
 'respawn':(400,330,260,140,'switch_stage respawn\nhealth minus one','data'),
 'reload':(770,330,260,140,'Rebuild same Level\nfrom current_level','primary'),
 'check':(1120,330,250,140,'Health check\nhealth <= 0 ?','ui'),
 'play':(1120,100,250,120,'Continue play\nhealth above zero','primary'),
 'menu':(1120,610,250,130,'Game over menu\nhealth at zero\nfreeze stage','danger'),
 'retry':(760,610,280,130,'Retry with Enter or R\nfull health','ui')}, [
 ('fall','respawn',''),('respawn','reload',''),('reload','check',''),('check','play','no'),('check','menu','yes'),('menu','retry','Enter / R'),('retry','reload','same stage')],
 'Retry preserves current_level and unlocked_level, resets health to five, and starts that stage at its TMX spawn point.')

DIAGRAMS['assets'] = diagram('diagram_assets', 'Asset loading path', {
 'game':(70,330,240,140,'Game.import_assets()','primary'),
 'helpers':(390,330,260,140,'support.py\nimage and folder helpers','data'),
 'root':(740,110,260,140,'PROJECT_ROOT\nfrom settings.py','ui'),
 'graphics':(740,340,260,140,'graphics/\nframes fonts maps','external'),
 'audio':(740,570,260,140,'audio/\nSound and music','external'),
 'cache':(1100,330,270,140,'Frame dictionaries\nUI and audio assets','data'),
 'stage':(1450,330,140,140,'Level\nuses assets','primary')}, [
 ('game','helpers','calls'),('helpers','root','resolve'),('root','graphics','path'),('root','audio','path'),('graphics','cache','load'),('audio','cache','load'),('cache','stage','pass')],
 'Paths are rooted at the project directory, so startup does not depend on the current shell folder.')

PAGES = [
('Executive summary', [
 'Tidebound is a single-player, side-scrolling platform game written in Python with Pygame. The completed implementation combines a tile-based stage, a node-and-path overworld, animated entities, hazards, collectibles, health, audio, and a small heads-up display. TMX files provide level geometry and object placement; Python classes turn those authored records into live sprites and collision groups.',
 'The central Game object owns the display, clock, loaded assets, session Data, maps, and the active stage. A frame updates the selected stage, updates the UI, checks health, and presents the display. Level and Overworld implement different interaction and rendering rules while sharing the same session data.',
 'This document describes the implementation present in the Tidebound workspace. It distinguishes observed behavior from suggested improvements. The current runtime uses a 1280 by 720 display, a 60 FPS frame cap, Python 3.10 or later syntax, Pygame, and PyTMX. The project is an offline local game with no network service or persistent save file.'
]),
('Purpose scope and evidence', [
 'The document is for developers who need to run, explain, maintain, or extend the game. It explains the repository structure, runtime architecture, map contract, data flow, gameplay systems, failure behavior, diagrams, and verification approach. It is not a player walkthrough or a claim that every gameplay path has been tested on every operating system.',
 'The source of truth is the checked-out `code_complete` implementation and its TMX map files. `code_start` is a tutorial scaffold with a smaller API and should not be treated as the runnable game. The dependency file declares Pygame and PyTMX. The graphics, audio, and map folders are runtime inputs.',
 'Implementation facts in this document follow the inspected source. The verification checklist is a proposed repeatable release practice unless it explicitly reports a performed check. Runtime behavior can still vary with display drivers, audio devices, and user-created map data.'
]),
('Product profile and boundaries', [
 'The player begins at level zero. The overworld represents stages as nodes connected by route objects. Direction keys move the icon where a route is available and unlocked; Enter selects the current node and loads its level. In a level, the player moves, jumps, attacks, collects items, avoids hazards, and reaches a flag to return to the overworld.',
 'The implemented loop is local and single-user. Session values live only in memory. A restart from the game-over overlay rebuilds the current stage, restores health to five, resets coins, and preserves the current level and highest unlocked level. Falling below a stage costs one health and reconstructs the same level. Reaching zero health pauses gameplay at the overlay.',
 'The repository does not implement online play, cloud saves, accounts, a settings menu, a level editor, or automatic package building. PyTMX loads the authored maps; Pygame handles input, timing, drawing, sprites, and sound. These boundaries inform the architectural diagrams and the proposed extension points later in the document.'
]),
('System context', [
 'The player interacts through keyboard input and observes the game through a Pygame display. The Game loop receives input and elapsed time, delegates to one active stage, and draws the result. Stage data and visual assets are local files. PyTMX turns the TMX document into layers and objects that Level and Overworld consume.',
 'The Pygame mixer loads sound effects and background music. There is no request to a remote service during normal play. The Python process owns all mutable session state, including health, coins, current level, and unlocked level. Closing the window terminates the process rather than saving a session.',
 'This context boundary is useful when debugging: a missing asset is a local path or packaging issue; a malformed object is a content contract issue; and an input or display issue is usually inside the Pygame runtime or its driver.'
], 'context'),
('Runtime component architecture', [
 'Game is the composition root. It initializes Pygame, creates the display and clock, imports assets, loads six numbered TMX stages and the overworld map, creates UI and Data, and chooses the initial Level. It also owns the stage-switch callback and the game-over menu.',
 'Level assembles scenery, terrain, player, moving objects, enemies, water, and collectible items from one map. It keeps groups for solid collision, one-way platforms, damage, enemies, projectiles, and items. Overworld builds nodes and route sprites and moves an icon between selectable stages. These stage classes are alternatives held by `Game.current_stage`.',
 'Player and enemy classes encapsulate movement or behavior. Sprite subclasses carry images, rectangles, z-order values, and update routines. AllSprites and WorldSprites specialize camera and draw ordering. UI reads session values through Data and maintains animated hearts and temporary coin display.'
], 'architecture'),
('Frame lifecycle', [
 'At the top of each iteration, Game asks the Pygame clock for elapsed milliseconds and caps the loop at 60 frames per second. It converts the tick duration to seconds and passes `dt` to the stage and UI. This makes movement and animation rates independent of the nominal frame rate, subject to the chosen collision and rectangle precision.',
 'Window-close events quit the process. While the game-over flag is active, gameplay updates stop; Enter or R rebuilds the current level and Escape quits. During play, the stage updates its sprite group, resolves collisions and constraints, and draws the world. The UI then draws health and coin indicators. If health is zero, Game draws the overlay before flipping the completed frame to the display.',
 'The loop is intentionally synchronous. Asset loading, TMX parsing, simulation, rendering, and input all execute on the main thread. That is appropriate for this project size and avoids coordination between worker threads and Pygame surfaces.'
], 'loop'),
('Stage state model', [
 'The runtime has three visible gameplay conditions: overworld, level, and game over. Startup constructs Level zero directly. Entering a node switches to a Level. Touching a level flag returns to the overworld and may increase `unlocked_level`. Falling below the bottom boundary reloads the same level and decreases health. Zero health changes the presentation to the game-over overlay.',
 'There is no separate title screen or pause state. The game-over overlay is a state of Game rather than a stage class. It freezes the active stage by skipping updates, while preserving the last rendered scene behind a dark translucent surface. The retry action retains the selected level and progression but creates new Data and UI instances.',
 'The switch callback uses a target string and, for successful exits, an unlock value. This is simple for the current set of transitions. If more stages or outcomes are added, a small enum or explicit transition object would reduce the risk of misspelled target names.'
], 'states'),
('Controls and player interaction', [
 'In a level, Left and Right set horizontal movement. Space requests a jump. Down activates the one-way platform skip timer. X requests an attack when the attack cooldown is inactive. Holding Space can repeatedly request jumps when the player is grounded because the implementation reads the current key state each frame rather than only key-down events.',
 'In the overworld, arrow keys request movement along the selected node’s available route. The icon follows the map-authored path points. Enter writes the selected node’s stage number to `Data.current_level` and invokes the stage switch. The current node is updated by sprite overlap after the icon moves.',
 'The game-over overlay accepts Enter or R to retry the same stage and Escape to quit. Window close also quits. Keyboard input is polled through `pygame.key.get_pressed()` for movement; discrete window and retry actions use the event queue.'
]),
('Session data model', [
 'Data owns four session values: `_health`, `_coins`, `current_level`, and `unlocked_level`. Health starts at five, coins at zero, and both level values at zero. The fields have no disk serialization, so a new process begins a new run.',
 'The coin property setter stores the requested amount, exchanges each full hundred coins for one health point, and asks UI to show the remaining coin count. The health setter updates the value and recreates the heart sprites. Item objects update coins for gold, silver, diamond, or skull pickups, and increase health for a potion.',
 'Game creates Data and passes the same object into the active stage and UI path. This shared reference keeps effects visible without a separate database or event bus. The tradeoff is tighter coupling: item and player logic can mutate session state directly. A future event model could make such changes easier to audit, but is not required for the current offline game.'
], 'model'),
('TMX level construction', [
 'Each stage map is parsed by `pytmx.util_pygame.load_pygame`. Level reads its dimensions and the custom properties in the Data layer, including background selection, level unlock value, top limit, and horizon line. It creates the camera-aware AllSprites group from that metadata before building the level entities.',
 'Tile layers named BG, Terrain, FG, and Platforms are iterated as `(x, y, surface)` tiles. Terrain tiles join the solid collision group, while Platforms tiles join the one-way platform group. Background and foreground tiles receive z values according to the current implementation’s layer mapping.',
 'Object layers then create background details, player and scenery, moving hazards and platforms, enemies, items, and water. Custom object properties configure values such as speed, direction, platform behavior, spike radius, and level unlock. A renamed layer or property can therefore cause a runtime failure; map edits must preserve the expected names and types.'
], 'map_build'),
('Player movement model', [
 'Player stores a visual rectangle and a smaller hitbox rectangle. The hitbox is derived by shrinking the image rectangle by 76 pixels horizontally and 36 vertically. Movement uses a direction vector, a horizontal speed of 200 units per second, gravity of 1300 units per second squared, and a jump impulse of 900 units per second upward.',
 'Each update snapshots the old hitbox, advances timers, reads input, applies horizontal movement, applies vertical movement and gravity, resolves solid and one-way collisions, carries the player with a moving platform when applicable, checks nearby floor and wall contacts, chooses an animation state, and applies flicker after damage.',
 'Wall sliding reduces vertical acceleration while the player touches a wall and the wall-slide block timer is inactive. Wall jumps set a horizontal impulse away from the contacted side. All movement uses elapsed time, while collision resolution compares old and current rectangles to infer which face crossed a surface.'
]),
('Collision and contact pipeline', [
 'Solid collision iterates the collision sprite group and checks rectangle overlap. Horizontal resolution compares the old hitbox with each solid sprite’s old rectangle to determine whether the player crossed a left or right boundary. Vertical resolution applies the same crossing approach to ceilings and floors and zeroes the vertical velocity after contact.',
 'One-way platform collision is a separate pass. It only blocks downward crossings from above and is disabled while the platform-skip timer is active. A short floor probe and side probes establish floor, left-wall, and right-wall contact for animation and wall movement decisions. Moving platforms are detected near the floor probe and carry the player by their own direction and speed.',
 'After sprite updates, Level checks projectile contacts with solid objects, hazard overlaps with the player hitbox, collectible overlaps with the player, and player attacks against Tooth and Pearl targets. The current loops are straightforward group scans. For larger maps, spatial hashing or pygame sprite collision helpers could reduce work, but should be measured before adding complexity.'
], 'collision'),
('Enemy behavior', [
 'Tooth is a ground patrol enemy. It chooses an initial direction, advances horizontally at 200 units per second, animates its run frames, and reverses when a floor probe detects a ledge or a wall probe detects an obstacle. Its hit timer limits how often an attack can reverse its direction.',
 'Shell is a stationary ranged enemy. It checks player distance, whether the player lies in front of its firing direction, and whether the vertical separation is below 30 units. If all conditions hold and its 3000 millisecond timer is inactive, it plays a fire animation and creates one Pearl at animation frame three.',
 'Enemies are created from the Enemies TMX layer. Tooth uses solid collision rectangles copied at construction. Shell receives a reference to the player and a callback for projectile creation. The current behavior is deterministic apart from Tooth’s initial direction and animation details; there is no pathfinding or enemy difficulty scaling.'
]),
('Projectile and attack behavior', [
 'Pearl is a moving projectile created by Shell. It begins offset from the shell by 50 pixels in its travel direction, moves at 150 units per second, and activates a five-second lifetime timer. A short reverse timer prevents repeated attack reversals in quick succession.',
 'The player’s attack state is driven by the X key, an attack cooldown timer, and animation frames. The Level attack pass checks whether the player is attacking, overlaps a Tooth or Pearl, and is facing toward that target. A valid contact calls the target’s reverse method rather than removing it.',
 'A projectile that hits a solid collision sprite is removed and creates a particle effect. A projectile that overlaps the player reduces health through the Player damage timer, is killed, and creates a particle. Hazard and projectile contact share the player’s damage cooldown, so multiple simultaneous overlaps do not repeatedly decrement health during the same brief invulnerability interval.'
]),
('Items health and coins', [
 'Item sprites are animated and centered at their map object positions. Gold adds five coins, silver adds one, diamond adds twenty, and skull adds fifty. A potion adds one health. Collection removes the first overlapping item, activates its effect, creates a particle animation, and plays the coin sound.',
 'Every full hundred coins converts into one health point. The setter subtracts one hundred and increments health once per assignment. As implemented, a single item cannot exceed 100 coins, so this conversion is adequate for current pickups. If larger rewards are introduced, the conversion should use a loop or quotient arithmetic so multiple lives are awarded correctly.',
 'Health is represented as an integer and displayed through a row of animated heart sprites. Health can decrease from hazards, enemy projectiles, and falling below the map. A fall triggers a same-level reconstruction and costs one health; zero health activates the game-over presentation.'
]),
('Overworld route selection', [
 'Overworld parses its tile layers, animated water, decorative objects, Paths, and Nodes. Each path object supplies start and end stage identifiers and a point list. Each node object supplies a stage number and directional properties that reference route identifiers, with an optional reverse marker.',
 'Node.can_move allows a direction only when the node declares that route and the destination stage number does not exceed `Data.unlocked_level`. The icon follows the route’s ordered points at 400 units per second. On arrival, sprite collision updates the current node; Enter then loads that node’s stage.',
 'PathSprites are synthesized from route points and node grid positions. The algorithm identifies vertical, horizontal, and corner tiles to select the matching path graphic. Route direction and map topology are thus content data rather than hard-coded navigation. When editing the overworld, node stage IDs, path IDs, and route endpoint properties must remain consistent.'
], 'world'),
('Rendering camera and depth', [
 'Level rendering uses AllSprites. The group stores world dimensions, camera borders, optional tiled background, cloud data, and horizon information. On draw, it centers the camera around the target and clamps the offset to the map limits. Sky scenes draw a background fill, sea, horizon line, large cloud strip, and sprites.',
 'Sprites carry a `z` value. AllSprites sorts them by z before blitting, allowing background tiles, clouds, scenery, main entities, water, and foreground effects to appear in a stable order. The overworld uses WorldSprites, centers its camera around the icon, draws background layers first, and sorts main sprites by vertical position.',
 'The display surface is cleared or painted by the stage, then the UI draws hearts and coin text over the world. The main loop flips the display once per frame. This arrangement makes camera behavior a responsibility of the sprite groups rather than each entity.'
], 'render'),
('Animation sprite and effects', [
 'AnimatedSprite stores an ordered frame list and a floating frame index. Each update advances the index by animation speed multiplied by elapsed seconds, then selects a frame modulo the frame count. This supports looping effects such as water, candles, palms, and player movement.',
 'ParticleEffectSprite uses the same frame progression but kills itself after the last frame. Heart animations are triggered randomly and return to an inactive state when their finite sequence completes. Player animation state is selected from idle, run, jump, fall, wall, attack, and air-attack states, with horizontal flipping when facing left.',
 'The asset importer sorts frame filenames numerically for folder-based animations and builds dictionaries from filename stems for keyed assets. Animation folder conventions are therefore part of the runtime contract. Missing or empty frame folders fail during entity construction rather than producing a fallback frame.'
]),
('Audio and timers', [
 'Game loads coin, attack, jump, damage, and Pearl sound effects into a dictionary. It loads the background music as a Sound object, sets its volume to 0.5, and loops playback continuously. Level adjusts coin and damage volume locally and plays effects at the point where interactions occur.',
 'Timer stores a duration, optional callback, start timestamp, active flag, and repeat behavior. It uses Pygame’s millisecond tick count. Player timers cover wall jump, wall-slide blocking, platform skipping, attack cooldown, and damage invulnerability. Shell uses a firing interval; Pearl uses lifetime and reverse intervals; AllSprites uses a repeating timer to create clouds.',
 'Timers are updated by the owning object. This keeps timing logic near the behavior it controls, but requires every active timer to receive update calls. A future central scheduler is unnecessary unless timer count and cross-system coordination grow substantially.'
]),
('Interface and game-over menu', [
 'UI draws animated hearts at the upper-left and temporarily displays coin count and a coin icon after a coin change. A one-second timer controls coin visibility. Health changes rebuild the heart sprite group to match the current value. The UI is rendered after the active stage so it remains above world art.',
 'Game checks health after a stage update. At zero or below, it sets `game_over` and draws a full-screen translucent overlay with a centered panel. The stage stops updating while the overlay is active, leaving the last scene visible beneath the menu.',
 'Enter or R retries the current level, restores health to five, resets coins, and preserves the current level and highest unlocked level. Escape or a window-close event quits. There is no mouse selection, options screen, or persisted best score in the current implementation.'
]),
('Death respawn and retry sequence', [
 'Falling is detected when the player hitbox bottom passes the level’s pixel height. Level calls the `respawn` transition and returns from its constraint check. Game decrements health, reconstructs a Level using the same `Data.current_level`, and places a new Player at the spawn object defined in that TMX map.',
 'The game loop checks health after the stage runs. If the fall used the last health, the same-level reload is created and the game-over overlay is shown before the next update. This preserves the level context during the menu. Hazard damage can also reduce health to zero while the player stays in the current stage; the same game-over behavior follows.',
 'Retry captures the selected level and unlocked progress before creating new UI and Data objects. It then restores those progress values, constructs the same Level, and clears the game-over flag. The session remains in memory only; closing the process discards it.'
], 'death'),
('Asset loading and path resolution', [
 'The asset root is derived from the location of `code_complete/settings.py`. `asset_path` joins a sequence of path components to the repository root and removes a leading parent marker for older import calls. This lets the game launch from a shell whose current directory is outside the repository.',
 'Support helpers load single images, ordered animation folders, keyed image dictionaries, and immediate subfolders. Pygame image conversion uses alpha for sprites and optional opaque conversion for backgrounds. The display must already be initialized before converted surfaces are loaded, which is satisfied because Game calls `pygame.init()` before `import_assets()`.',
 'The code keeps loaded surfaces and audio objects in dictionaries for reuse by stages. Stages create sprites from those shared assets rather than reloading image files for each entity. That reduces repeated disk access and keeps stage construction predictable.'
], 'assets'),
('Python Pygame and PyTMX compatibility', [
 'The implementation uses structural pattern matching in Level.setup, so Python 3.10 or newer is required. The declared dependency bounds request Pygame 2.6.1 or newer while remaining below major version 3, and PyTMX 3.32 or newer below major version 4.',
 'Some Pygame builds expose floating rectangles through `Surface.get_frect` and `pygame.FRect`; others expose only integer Rect APIs. The settings module selects the floating API where available and provides a Rect fallback. Float rectangles preserve subpixel movement more accurately; the fallback prioritizes startup compatibility but quantizes coordinates to integers.',
 'The runtime also relies on mixer support for WAV and MP3 decoding, image conversion after display initialization, Pygame sprite groups, vector arithmetic, and PyTMX’s Pygame loader. A clean install should use a virtual environment and install from requirements.txt so the game dependencies do not conflict with unrelated Python applications.'
]),
('Tuning parameters and configuration', [
 'Global presentation parameters live in settings.py: window width and height are 1280 by 720, tile size is 64 pixels, and baseline animation speed is six frames per second multiplier. Z-layer constants define the broad drawing order. Game’s run method caps its clock at 60 frames per second.',
 'Player movement values are local to Player: horizontal speed 200, gravity 1300, jump impulse 900, and a hitbox inset from the sprite image. Enemy and projectile values live in their classes or map properties. TMX values configure patrol speed, platform status, flipping, spike radius and angular bounds, background tile selection, and unlock stage.',
 'Keeping configuration close to behavior makes the tutorial code easy to follow. For a larger project, a validated configuration object could collect tunable values and report missing or invalid map properties at startup. Any configuration refactor should preserve the values currently authored in map objects.'
]),
('Repository and module inventory', [
 '`code_complete` contains the playable implementation. `main.py` coordinates startup and the game loop; `level.py` builds and runs a stage; `overworld.py` handles node selection; `player.py` implements the player; `sprites.py` contains common sprite, item, effect, platform, node, icon, and path classes.',
 '`enemies.py` implements Tooth, Shell, and Pearl. `groups.py` contains camera-aware level and overworld groups. `ui.py` contains health and coin display. `data.py` owns session values. `timer.py` provides elapsed-time timers. `settings.py` defines shared constants and asset paths. `support.py` loads art. `debug.py` is a small display helper.',
 '`code_start` is a reduced tutorial starting point and does not implement the complete game systems. `data/levels` and `data/overworld` contain TMX maps; `data/tilesets` contains tileset definitions; `graphics` and `audio` contain runtime assets. requirements.txt declares the Python packages.'
]),
('Core class contracts', [
 'Game constructs Level with a parsed map, level frame dictionary, audio dictionary, Data object, and stage-switch callback. Overworld receives its map, Data, overworld assets, and the same callback. Both stages expose `run(dt)` and own the sprites they create.',
 'Player receives a position, sprite groups, collision groups, player frames, Data, attack sound, and jump sound. Level creates it before enemy objects that need a player reference. Enemy projectiles call Level.create_pearl through a callback, keeping the projectile constructor independent of stage internals.',
 'Sprite constructors generally register themselves into one or more Pygame groups. Group membership drives updates, collision passes, and drawing. Because registration often occurs during construction, constructors require a valid display surface, loaded assets, and initialized groups. These implicit preconditions should be retained or made explicit if classes are reused outside Game.'
]),
('Startup and run operations', [
 'From the repository root, create and activate a virtual environment, install requirements.txt, then run `python code_complete/main.py`. Asset paths resolve from the source file, so launching from another working directory is supported. The game opens a 1280 by 720 window titled Tidebound and starts the background track.',
 'Pygame initialization happens before font, image, and sound loading. The TMX files are then loaded into a dictionary keyed from zero through five, along with the overworld map. Data begins at level zero, and Game constructs that first Level before starting the music and frame loop.',
 'To exit during gameplay, close the window. At game over, press Escape or close the window. To continue after depletion, press Enter or R to retry. For a headless smoke check, use Pygame dummy video and audio drivers and instantiate Game; interactive controls still require a real window.'
]),
('Diagnostics and developer workflow', [
 'When startup fails, first confirm that dependencies installed in the same interpreter used to launch the game. A `ModuleNotFoundError` points to environment setup. A file-not-found error usually points to an absent asset, incorrect asset name, or incomplete repository checkout. An image conversion failure can indicate initialization order or a damaged image.',
 'When TMX loading succeeds but stage construction fails, inspect the referenced layer and object names in the map, then compare required custom properties with the Level or Overworld constructor. Key lookups such as `get_layer_by_name` and dictionary indexing deliberately fail loudly when content contracts are broken.',
 'For movement bugs, check the player’s hitbox and old rectangle values, the relevant collision group, and `dt`. For progression bugs, inspect node stage IDs, route endpoints, directional route properties, and `Data.unlocked_level`. The debug helper can draw a value to the display but should be removed from production paths if it interferes with layout.'
]),
('Verification matrix', [
 'A release check should cover startup from the repository root and an unrelated working directory; loading all six numbered maps and the overworld; one update and draw pass per stage; movement, jumping, wall interaction, and platform dropping; and item, hazard, projectile, and flag interactions.',
 'Progression checks should confirm that locked nodes remain unavailable, successful completion raises the unlock value monotonically, and the selected node loads the matching map. Failure checks should confirm that falling costs one health, rebuilds the same stage at its spawn, and zero health pauses gameplay at the overlay. Retry should preserve current level and unlock progress while restoring health and resetting coins.',
 'The workspace checks performed for this document include Python compilation, Game initialization under dummy video and audio drivers, one frame through each of six maps and the overworld, same-stage respawn, game-over drawing, and retry state reset. These checks do not replace interactive play testing or cover every frame timing and display driver.'
]),
('Performance profile', [
 'The project is a small 2D game with sprite counts expected to remain modest. Pygame draws surfaces through blits and the game loop has a 60 FPS cap. Images and sounds load once during initialization and are reused. Animation changes a sprite’s surface reference rather than rebuilding assets each frame.',
 'The most visible scaling costs are full sprite-group update and draw passes, repeated sorting for z-order, and player collision checks that iterate groups and rebuild probe rectangle lists. These operations are simple and readable. Optimizing them without an observed frame-time problem would add complexity without a measured benefit.',
 'If levels grow, profile representative scenes first. Likely options include caching stable draw order, culling sprites outside the camera, reusing probe structures, and indexing collision candidates by tile or spatial cell. Any optimization must preserve visual order and moving-platform collision behavior.'
]),
('Reliability and failure handling', [
 'The current program relies on Python exceptions for missing required data. This is suitable during development because it identifies the bad lookup. A packaged game could catch startup failures at the top level and show a useful message naming the asset or layer that could not load.',
 'TMX authoring is trusted input. A missing player object can leave `self.player` unset; a missing flag can leave the finish rectangle unset; missing properties can trigger key errors; and an empty animation folder can cause an index error. The shipped maps satisfy these contracts in the inspected run, but the classes do not validate them explicitly.',
 'The game-over menu and same-stage respawn provide recovery from gameplay failure. There is no crash-reporting, autosave, or recovery after process failure. A future validation command could scan every map at startup and report missing layers, required objects, and properties before play begins.'
]),
('Privacy and security posture', [
 'The application operates as a local game. The inspected runtime does not create user accounts, transmit gameplay data, call web APIs, or write save files. Session health, coins, and progression exist only in memory until the process ends.',
 'The principal trust boundary is the local project directory: TMX files, images, fonts, and audio are loaded from it. The project uses PyTMX and Pygame decoders, so distributable builds should include only assets from trusted sources and supported file types. The game should not be run with elevated permissions because it does not require them.',
 'A packaged release should document its third-party dependency versions and asset licensing. The original repository describes the tutorial code as CC0; that statement does not automatically establish the license for every third-party asset included in a redistributed build. Confirm asset rights before publishing.'
]),
('Accessibility and usability review', [
 'Keyboard controls cover movement, level selection, retry, and quit. The screen uses a fixed 1280 by 720 window. Health is shown through heart graphics, while coins use both text and an icon. There is no rebinding, controller support, scaling menu, audio slider, subtitles, or color-blind mode.',
 'The game-over menu communicates the available keys in text and uses a high-contrast panel over the frozen scene. Important status is not conveyed only by color: heart count is visual and coin count is numeric. The menu text is rendered with Pygame’s default font, while gameplay UI uses the project font.',
 'Useful next steps include separating key mapping from gameplay logic, adding window scaling that preserves aspect ratio, providing independent music and effect volume controls, and testing text size at common display resolutions. These are recommendations, not features already present.'
]),
('Packaging and release procedure', [
 'A source release needs the Python files, requirements.txt, maps, tileset definitions, graphics, and audio. Running only the Python modules without the content directories will fail during asset loading. The `code_start` tutorial scaffold is optional for a player-focused release but is part of the repository’s teaching context.',
 'A release candidate should be installed into a clean virtual environment, launched from outside the repository, and smoke-tested with every level map. Verify that audio decodes on the target platform and that the included font and image paths preserve case and spaces. Windows and case-sensitive filesystems differ in path handling.',
 'If producing a standalone executable, use a packaging tool only after checking how it collects dynamic asset directories and PyTMX map dependencies. Bundle assets in a known resource directory and derive runtime paths from the packaged application location. The repository currently has no build script or executable configuration.'
]),
('TMX authoring guide', [
 'Level maps depend on named layers: Data, BG, Terrain, FG, Platforms, BG details, Objects, Moving Objects, Enemies, Items, and Water. Level construction queries these names directly. Keep the expected layer present even when it contains no objects if the loader expects it.',
 'Object names select constructors and frame dictionaries. The player object creates Player; flags define the finish rectangle; barrels and crates become solid sprites; saws and floor spikes become damage sources; moving objects require speed, flip, and platform properties; spike objects require radius and angle metadata; shells require reverse; items use item names as keys.',
 'The overworld requires tile layers named main and top, plus Objects, Paths, and Nodes. Node properties define stage and directional routes. Path objects declare start and end stage IDs and ordered points. Since map data is executable configuration, every change should be validated by loading the map and constructing its stage.'
]),
('Technical debt and constraints', [
 'The code uses wildcard imports from settings and support, which makes dependencies less explicit. Several methods depend on string layer names and dynamic object names. A rename can therefore fail at runtime without static type support. Typed data structures and explicit imports would improve navigation as the project grows.',
 'Data setters update UI directly, coupling model and presentation. Some collision logic scans all sprites and sorts all sprites each frame. The compatibility fallback from floating rectangles to integer rectangles can reduce subpixel movement precision in standard Pygame builds. TMX properties are accessed without a schema validation pass.',
 'The game has no automated regression suite, persistent saves, logging strategy, packaging configuration, or user settings. These are manageable constraints for a tutorial-scale game. Prioritize changes by observed failure or a concrete feature requirement rather than broad refactoring.'
]),
('Extension roadmap', [
 'A low-risk next step is a map validator that checks required layers, objects, and custom properties for every TMX file before launch. It can report the map name and missing field without changing gameplay. A second step is a small regression suite for data setters, route unlocks, stage transitions, and retry semantics.',
 'A content expansion can introduce a settings panel for audio and controls, a persistent save file for progression, and support for resolution scaling. These require decisions about save format migration, key binding representation, and aspect-ratio behavior before implementation.',
 'A performance step should begin with frame-time profiling on the busiest level. Cache draw order or add camera culling only if measurements show a meaningful cost. An architecture step can replace string transition targets with explicit transition types and decouple UI notifications from Data setters.'
]),
('Requirements traceability', [
 'The runtime requirement to load gameplay content is implemented by Game’s asset and map loading and by Level.setup. Player movement and combat requirements are handled by Player, Level collision passes, Tooth, Shell, and Pearl. Overworld navigation is implemented by Node, Icon, PathSprite, and Overworld.',
 'Progression and failure requirements are carried by Data and Game.switch_stage. Successful flags use the level’s `level_unlock` property; a fall uses the `respawn` target and reloads `current_level`; health at or below zero activates the game-over overlay; retry restores health and reconstructs that same level while preserving unlock progress.',
 'Visual requirements are handled by AllSprites, WorldSprites, z-layer constants, and UI. Runtime timing uses Clock and Timer. This mapping helps reviewers locate the code that owns a behavior before changing it and identifies areas where a future automated test can provide direct coverage.'
]),
('Glossary and map inventory', [
 'A stage is one playable TMX level. The overworld is the route-selection map. A node is a selectable stage marker. A path is a sequence of points connecting nodes. A sprite group is a Pygame collection used for update, draw, or collision membership. A hitbox is the player’s reduced collision rectangle. A z layer is a draw-order value.',
 'The level directory contains numbered files `0.tmx` through `5.tmx`, plus `omni.tmx`. The running Game loads the six numbered maps into the playable map dictionary. The overworld map is loaded separately. `omni.tmx` remains in the repository but is not included in the current six-stage dictionary.',
 'Tiles are authored on a 64-pixel grid. Map object coordinates and dimensions are converted into world positions using those tile values. The camera centers on the player or overworld icon and clamps to the world boundaries calculated from map dimensions.'
]),
('Reference tables and final checklist', [
 'Key runtime values are: window 1280 by 720 pixels; tile size 64 pixels; frame cap 60 FPS; initial health five; initial coins zero; player horizontal speed 200; gravity 1300; jump impulse 900; Tooth speed 200; Pearl speed 150; Shell fire interval 3000 milliseconds; Pearl lifetime 5000 milliseconds; and hit invulnerability 400 milliseconds.',
 'Before publishing a change, install declared dependencies in a clean environment, compile Python files, launch from outside the repository, instantiate all numbered maps and the overworld, exercise a representative interaction, and inspect the game-over and retry paths. For map changes, validate all required layers and object properties. For asset changes, confirm filenames, image conversion, and audio decoding.',
 'The architecture is deliberately compact: Game coordinates, stage classes own play modes, sprite classes own entity behavior, groups own drawing, Data owns session values, and TMX files own authored layout. Keeping those responsibilities clear is the safest basis for adding content or improving the game.'
])
]

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill.lstrip('#'))
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    val = OxmlElement('w:t'); val.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for e in (begin,instr,sep,val,end): run._r.append(e)

def add_paragraph(doc, text, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    p.add_run(text)
    return p

def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0,0)
    set_cell_shading(cell, 'EEF3F7')
    set_cell_margins(cell, 130, 180, 130, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper() + '  ')
    r.bold = True; r.font.color.rgb = RGBColor(46,111,149)
    p.add_run(text)

def add_toc_page(doc):
    doc.add_page_break()
    h = doc.add_heading('How to use this document', 1)
    h.paragraph_format.space_after = Pt(10)
    add_paragraph(doc, 'This technical reference describes the completed Tidebound game and its runtime design. The diagrams use the same module names as the source code. Current behavior is described as implemented; improvement ideas are identified as recommendations.')
    groups = [
      ('System design', ['Executive summary','Purpose scope and evidence','Product profile and boundaries','System context','Runtime component architecture','Frame lifecycle','Stage state model']),
      ('Gameplay systems', ['Controls and player interaction','Session data model','TMX level construction','Player movement model','Collision and contact pipeline','Enemy behavior','Projectile and attack behavior','Items health and coins','Overworld route selection']),
      ('Runtime engineering', ['Rendering camera and depth','Animation sprite and effects','Audio and timers','Interface and game-over menu','Death respawn and retry sequence','Asset loading and path resolution','Python Pygame and PyTMX compatibility','Tuning parameters and configuration']),
      ('Operations and reference', ['Repository and module inventory','Core class contracts','Startup and run operations','Diagnostics and developer workflow','Verification matrix','Performance profile','Reliability and failure handling','Privacy and security posture','Accessibility and usability review','Packaging and release procedure','TMX authoring guide','Technical debt and constraints','Extension roadmap','Requirements traceability','Glossary and map inventory','Reference tables and final checklist'])]
    for title, items in groups:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title); r.bold = True; r.font.color.rgb = RGBColor(0,0,0)
        p2 = doc.add_paragraph('  •  '.join(items))
        p2.paragraph_format.space_after = Pt(5)
        p2.paragraph_format.line_spacing = 1.0

def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.35)
    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(30,42,54)
    normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [('Title',32,NAVY),('Heading 1',22,NAVY),('Heading 2',15,BLUE),('Heading 3',12,TEAL)]:
        st = styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.keep_with_next=True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = header.add_run('TIDEBOUND  /  TECHNICAL ARCHITECTURE')
    rr.font.name='Arial'; rr.font.size=Pt(8); rr.font.bold=True; rr.font.color.rgb=RGBColor(0,0,0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr=footer.add_run('Tidebound Technical Reference     •     '); fr.font.size=Pt(8); fr.font.color.rgb=RGBColor(88,104,118)
    add_page_field(footer)
    # Cover
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(115); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('TIDEBOUND'); r.font.name='Arial'; r.font.size=Pt(18); r.font.bold=True; r.font.color.rgb=RGBColor(46,111,149)
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(16)
    p.add_run('Technical Architecture\nand Design Reference')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('A developer guide to the Pygame platform game'); r.font.size=Pt(15); r.font.color.rgb=RGBColor(0,0,0)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(24)
    r=p.add_run('Runtime design  •  Gameplay systems  •  Software diagrams'); r.font.size=Pt(11); r.font.color.rgb=RGBColor(46,111,149)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(105)
    r=p.add_run('Prepared 16 October 2026\nBased on the Tidebound workspace source'); r.font.size=Pt(10); r.font.color.rgb=RGBColor(88,104,118)
    add_toc_page(doc)
    for idx, page in enumerate(PAGES):
        title, paragraphs, *rest = page
        diagram_key = rest[0] if rest else None
        doc.add_page_break()
        h = doc.add_heading(title, 1)
        h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(10)
        for paragraph in paragraphs:
            add_paragraph(doc, paragraph)
        if diagram_key:
            dpath = DIAGRAMS[diagram_key]
            # Diagrams remain large enough to read while leaving room for the caption.
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2)
            shape = p.add_run().add_picture(str(dpath), width=Inches(6.85))
            shape._inline.docPr.set('descr', f'Software diagram showing {title.lower()} and its main components and transitions.')
            cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(4)
            cr=cap.add_run(f'Figure {sum(1 for x in PAGES[:idx+1] if len(x)>3)}  {title}')
            cr.italic=True; cr.font.size=Pt(8.5); cr.font.color.rgb=RGBColor(88,104,118)
        else:
            note = 'Implementation reference'
            callout = paragraphs[-1][:235] if paragraphs else ''
            add_callout(doc, note, callout)
    # Document metadata
    doc.core_properties.title = 'Tidebound Technical Architecture and Design Reference'
    doc.core_properties.subject = 'Software architecture, gameplay systems, operations, and diagrams for Tidebound'
    doc.core_properties.author = 'OpenAI'
    doc.core_properties.keywords = 'Tidebound, Pygame, Python, architecture, technical documentation'
    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    print('content sections:', len(PAGES), 'diagrams:', len(DIAGRAMS))

if __name__ == '__main__':
    make_doc()
