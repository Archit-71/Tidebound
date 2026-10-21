from docx import Document
from pathlib import Path
from zipfile import ZipFile
root=Path(r'C:\Users\Archit\Documents\Codex\2026-10-14\ok\outputs')
files=['Tidebound_Project_Report.docx','Tidebound_Technical_Architecture_and_Design.docx','Tidebound_Combined_Project_and_Technical_Document.docx']
terms=['github','super-pirate-world','super pirate world','clear-code-projects','clear code projects','developer','openai']
for name in files:
 p=root/name; d=Document(p); hits=[]
 texts=[x.text for x in d.paragraphs]
 for t in d.tables:
  for r in t.rows:
   texts += [c.text for c in r.cells]
 for sec in d.sections: texts += [x.text for x in sec.header.paragraphs+sec.footer.paragraphs]
 for t in terms:
  found=[x for x in texts if t in x.lower()]
  if found: hits.append((t,found[:2]))
 with ZipFile(p) as z:
  raw=b' '.join(z.read(n) for n in z.namelist() if n.endswith('.xml')).lower()
  hidden=[t for t in [b'github.com/clear-code-projects',b'super-pirate-world',b'openai'] if t in raw]
  print(name,'author=',repr(d.core_properties.author),'figures=',len(d.inline_shapes),'hits=',hits,'hidden-ids=',hidden,'zip=',z.testzip())
