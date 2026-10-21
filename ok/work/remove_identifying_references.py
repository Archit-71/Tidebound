from pathlib import Path
import re
from docx import Document

OUT = Path(r'C:\Users\Archit\Documents\Codex\2026-10-14\ok\outputs')
FILES = [
    OUT / 'Tidebound_Project_Report.docx',
    OUT / 'Tidebound_Technical_Architecture_and_Design.docx',
    OUT / 'Tidebound_Combined_Project_and_Technical_Document.docx',
]

REPLACEMENTS = [
    (r'A developer guide to the Pygame platform game', 'A project guide to the Pygame platform game'),
    (r'developer workflow', 'contributor workflow'),
    (r'developers who need', 'maintainers and contributors who need'),
    (r'Repository and module inventory', 'Project files and module inventory'),
    (r'repository structure', 'project file structure'),
    (r'from the repository root', 'from the project folder'),
    (r'repository root', 'project root'),
    (r'outside the repository', 'from another folder'),
    (r'the repository does not implement', 'the current build does not implement'),
    (r'the repository currently has', 'the current project has'),
    (r'the repository.s actual license', "the project's actual license"),
    (r'the repository.s', "the project's"),
    (r'upstream repository revision', 'upstream project revision'),
    (r'upstream repository', 'upstream project'),
    (r'by the repository', 'by the project'),
    (r'evidenced by the repository', 'evidenced by the project'),
]

def all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs

for path in FILES:
    doc = Document(path)
    for paragraph in all_paragraphs(doc):
        original = paragraph.text
        updated = original
        for pattern, replacement in REPLACEMENTS:
            updated = re.sub(pattern, replacement, updated, flags=re.IGNORECASE)
        # These affected paragraphs are single-style labels or prose. Keeping
        # the first run preserves paragraph styling while replacing the text.
        if updated != original:
            runs = paragraph.runs
            if runs:
                runs[0].text = updated
                for run in runs[1:]:
                    run.text = ''
            else:
                paragraph.add_run(updated)
    doc.core_properties.author = ''
    doc.save(path)
    print(path)
